from __future__ import annotations

import io
import ipaddress
import os
import shutil
import socket
import subprocess
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from urllib.parse import urldefrag, urljoin, urlsplit
from urllib.robotparser import RobotFileParser

import httpx
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .store import digest

PARSER_VERSION = "blocks-4"
AGENT = "ComputePolicyResearch/0.1"
ATTACHMENTS = {".pdf", ".docx", ".xlsx", ".txt", ".doc", ".xls", ".zip"}
MAX_BYTES = 15 * 1024 * 1024


class AcquisitionError(ValueError):
    pass


def resolve_public(host, port, **kwargs):
    """Optional DoH for networks whose system resolver returns proxy fake-IP ranges."""
    if os.getenv("POLICY_DNS_MODE", "system") == "system":
        return socket.getaddrinfo(host, port, **kwargs)
    if os.getenv("POLICY_DNS_MODE") != "doh":
        raise AcquisitionError("POLICY_DNS_MODE must be system or doh")
    try:
        address = ipaddress.ip_address(host)
        family = socket.AF_INET6 if address.version == 6 else socket.AF_INET
        return [(family, socket.SOCK_STREAM, 6, "", (str(address), port))]
    except ValueError:
        pass
    # Only a public hostname is sent to this fixed DNS provider, never policy/company contents.
    response = httpx.get(
        "https://dns.google/resolve",
        params={"name": host, "type": "A", "edns_client_subnet": "0.0.0.0/0"},
        timeout=15,
        trust_env=False,
    )
    response.raise_for_status()
    data = response.json()
    if data.get("Status") != 0 or data.get("TC"):
        raise AcquisitionError("public DNS did not return a complete successful answer")
    return [
        (socket.AF_INET, socket.SOCK_STREAM, 6, "", (item["data"], port))
        for item in data.get("Answer", [])
        if item["type"] == 1
    ]


@dataclass
class Resource:
    url: str
    body: bytes
    content_type: str
    blocks: list[dict] = field(default_factory=list)
    links: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    title: str = ""


def normalize(url: str) -> str:
    url = urldefrag(url.strip())[0]
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise AcquisitionError("only public HTTP(S) URLs without credentials are supported")
    if parsed.port not in {None, 80, 443}:
        raise AcquisitionError("non-standard ports are outside this crawler's scope")
    return str(httpx.URL(url))


class Fetcher:
    """Bounded fetcher. DNS results are validated and pinned for every connection and redirect."""

    def __init__(self, hosts: list[str], *, transport=None, resolver=resolve_public, delay=1.0):
        self.hosts = {host.lower() for host in hosts}
        self.transport = transport
        self.resolver = resolver
        self.delay = delay
        self.last_request = {}
        self.robots = {}
        self.client = httpx.Client(
            timeout=httpx.Timeout(25, connect=10),
            follow_redirects=False,
            trust_env=False,
            transport=transport,
        )

    def close(self):
        self.client.close()

    def _request(self, url, *, public_link=False):
        for attempt in range(2):
            try:
                result = self._request_once(url, public_link=public_link)
            except (httpx.RemoteProtocolError, httpx.ReadTimeout):
                if attempt:
                    raise
                time.sleep(1)
                continue
            if result[0] not in {429, 502, 503, 504} or attempt:
                return result
            time.sleep(1)

    def _request_once(self, url, *, public_link=False):
        url = normalize(url)
        host = urlsplit(url).hostname
        blocked = set(os.getenv("POLICY_BLOCKED_HOSTS", "").split(",")) - {""}
        if host in blocked:
            raise AcquisitionError("destination excluded from public research")
        if host not in self.hosts and not public_link:
            raise AcquisitionError(f"host outside source allowlist: {host}")
        port = 443 if urlsplit(url).scheme == "https" else 80
        addresses = sorted(
            {item[4][0] for item in self.resolver(host, port, type=socket.SOCK_STREAM)},
            key=lambda address: (ipaddress.ip_address(address).version != 4, address),
        )
        if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):
            raise AcquisitionError(
                "non-public destination blocked; 代理 fake-IP 网络可设置 POLICY_DNS_MODE=doh"
            )
        if any(address in blocked for address in addresses):
            raise AcquisitionError("destination excluded from public research")
        pause = self.delay - (time.monotonic() - self.last_request.get(host, 0))
        if pause > 0:
            time.sleep(pause)
        self.last_request[host] = time.monotonic()
        headers = {"User-Agent": AGENT, "Host": urlsplit(url).netloc, "Accept": "*/*"}
        # Prefer IPv4 on hosts without IPv6 routing; every fallback stays pinned to this validated set.
        attempts = addresses[:1] if self.transport else addresses[:3]
        for index, address in enumerate(attempts):
            target = httpx.URL(url) if self.transport else httpx.URL(url).copy_with(host=address)
            try:
                with self.client.stream(
                    "GET", target, headers=headers, extensions={"sni_hostname": host}
                ) as response:
                    if response.is_redirect:
                        return response.status_code, response.headers, b""
                    chunks, total = [], 0
                    for chunk in response.iter_bytes():
                        total += len(chunk)
                        if total > MAX_BYTES:
                            raise AcquisitionError("resource exceeds 15 MB limit")
                        chunks.append(chunk)
                    return response.status_code, response.headers, b"".join(chunks)
            except (httpx.ConnectError, httpx.ConnectTimeout):
                if index == len(attempts) - 1:
                    raise

    def _allowed(self, url, *, public_link=False):
        parsed = urlsplit(url)
        if parsed.hostname not in self.hosts and not public_link:
            raise AcquisitionError(f"host outside source allowlist: {parsed.hostname}")
        origin = f"{parsed.scheme}://{parsed.netloc}"
        if origin not in self.robots:
            robots_url = origin + "/robots.txt"
            for _ in range(6):
                status, headers, body = self._request(robots_url, public_link=True)
                if status in {301, 302, 303, 307, 308} and headers.get("location"):
                    robots_url = normalize(urljoin(robots_url, headers["location"]))
                    continue
                break
            else:
                raise AcquisitionError("robots.txt redirect limit exceeded")
            if status in {404, 410}:
                self.robots[origin] = None
            elif status == 200:
                if b"<html" in body[:1000].lower() or b"<!doctype html" in body[:1000].lower():
                    raise AcquisitionError("robots.txt returned HTML instead of crawl rules")
                parser = RobotFileParser()
                parser.parse(body.decode("utf-8", errors="replace").splitlines())
                self.robots[origin] = parser
            else:
                raise AcquisitionError(f"robots.txt unavailable (HTTP {status}); manual inspection needed")
        parser = self.robots[origin]
        if parser and not parser.can_fetch(AGENT, url):
            raise AcquisitionError("robots.txt disallows this path")
        if parser:
            crawl_delay = parser.crawl_delay(AGENT) or 0
            if crawl_delay > 30:
                raise AcquisitionError("publisher requires a crawl delay beyond this worker's budget")
            self.delay = max(self.delay, crawl_delay)

    def fetch(self, url) -> Resource:
        current = normalize(url)
        for _ in range(6):
            self._allowed(current)
            status, headers, body = self._request(current)
            if status in {301, 302, 303, 307, 308}:
                if "location" not in headers:
                    raise AcquisitionError("redirect has no destination")
                current = normalize(urljoin(current, headers["location"]))
                continue
            if status != 200:
                raise AcquisitionError(f"HTTP {status}; source content not verified")
            return parse(Resource(current, body, headers.get("content-type", "application/octet-stream")))
        raise AcquisitionError("too many redirects")

    def fetch_linked(self, url, parent: Resource) -> Resource:
        """Only a verified page's explicit attachment links may extend the host list."""
        if not any(link["url"] == url and link["attachment"] for link in parent.links):
            raise AcquisitionError("attachment was not linked by the source document")
        host = urlsplit(normalize(url)).hostname
        child = Fetcher(
            self.hosts | {host}, transport=self.transport, resolver=self.resolver, delay=self.delay
        )
        try:
            return child.fetch(url)
        finally:
            child.close()


def parse(resource: Resource) -> Resource:
    raw = resource.body
    suffix = PurePosixPath(urlsplit(resource.url).path).suffix.lower()
    rows = []
    mime = resource.content_type.lower()
    if (
        "html" in mime
        or suffix in {".html", ".htm"}
        or raw.lstrip().lower().startswith((b"<!doctype html", b"<html"))
    ):
        soup = BeautifulSoup(raw, "html.parser")
        resource.title = soup.title.get_text(" ", strip=True) if soup.title else resource.url
        # Government article metadata can include a withdrawal marker absent from
        # the site's generic <title> and stale search-engine snippets.
        article_title = soup.find("meta", attrs={"name": lambda v: v and v.lower() == "articletitle"})
        if article_title and article_title.get("content", "").strip():
            resource.title = article_title["content"].strip()[:500]
        for tag in soup.select("script,style,noscript,iframe"):
            tag.decompose()
        for anchor in soup.select("a[href]"):
            try:
                url = normalize(urljoin(resource.url, anchor["href"]))
            except (AcquisitionError, ValueError):
                continue
            text = anchor.get_text(" ", strip=True)
            ext = PurePosixPath(urlsplit(url).path).suffix.lower()
            resource.links.append(
                {
                    "url": url,
                    "text": text,
                    "attachment": ext in ATTACHMENTS
                    or (not ext and "附件" in text and not any(w in text for w in ["APP", "App", "客户端"])),
                }
            )
        # Prefer publisher article containers, retaining whole body only as a fallback.
        main = (
            soup.select_one("article, .TRS_Editor, #zoom, .article-content, .content-detail")
            or soup.body
            or soup
        )
        tags = main.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr"])
        for index, tag in enumerate(tags):
            if tag.find_parent(["p", "li", "tr"]):
                continue
            text = tag.get_text(" | " if tag.name == "tr" else " ", strip=True)
            if text:
                rows.append((f"html:{tag.name}:{index}", text))
        if not rows:
            rows = [
                (f"html:line:{index}", line.strip())
                for index, line in enumerate(main.get_text("\n", strip=True).splitlines())
                if line.strip()
            ]
        plain = main.get_text(" ", strip=True)
        if len(plain) < 100:
            resource.warnings.append("正文过短，可能需要浏览器渲染或人工导入")
        if any(
            word in plain for word in ["安全验证", "访问验证", "人机验证", "Access Denied", "Just a moment"]
        ):
            resource.warnings.append("疑似验证或拦截页面，不能作为完整政策证据")
    elif raw.startswith(b"%PDF") or suffix == ".pdf":
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > 150:
            raise AcquisitionError("PDF exceeds 150-page budget")
        for page_no, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if len(text.strip()) < 20:
                resource.warnings.append(f"第 {page_no} 页需 OCR/人工检查")
            for index, line in enumerate(text.splitlines()):
                if line.strip():
                    rows.append((f"pdf:page:{page_no}:line:{index}", line.strip()))
    elif suffix in {".docx", ".xlsx"} or "openxmlformats" in mime:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise AcquisitionError("Office archive expands beyond 50 MB")
        if suffix == ".xlsx" or "spreadsheetml" in mime:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                for sheet in workbook:
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > 200000:
                        raise AcquisitionError("spreadsheet declared dimensions exceed 200,000 cells")
                    for cells in sheet.iter_rows():
                        for cell in cells:
                            if cell.data_type == "f":
                                resource.warnings.append(
                                    f"{sheet.title}!{cell.coordinate} 为公式，需人工核对计算值"
                                )
                                continue
                            if cell.value is not None:
                                rows.append((f"xlsx:{sheet.title}:{cell.coordinate}", str(cell.value)))
                                if len(rows) > 20000:
                                    raise AcquisitionError("spreadsheet exceeds 20,000 populated cells")
            finally:
                workbook.close()
        else:
            document = Document(io.BytesIO(raw))
            rows.extend(
                (f"docx:paragraph:{i}", p.text) for i, p in enumerate(document.paragraphs) if p.text.strip()
            )
            for ti, table in enumerate(document.tables):
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        if cell.text.strip():
                            rows.append((f"docx:table:{ti}:row:{ri}:col:{ci}", cell.text))
    elif suffix == ".doc" and (shutil.which("textutil") or shutil.which("antiword")):
        try:
            converted = subprocess.run(
                (
                    [
                        "textutil",
                        "-convert",
                        "txt",
                        "-format",
                        "doc",
                        "-stdin",
                        "-stdout",
                        "-encoding",
                        "UTF-8",
                        "-noload",
                        "-nostore",
                    ]
                    if shutil.which("textutil")
                    else ["antiword", "-m", "UTF-8.txt", "-"]
                ),
                input=raw,
                capture_output=True,
                timeout=20,
                check=True,
            )
            text = converted.stdout.decode("utf-8")
            if len(text) > 500000:
                raise AcquisitionError("legacy DOC text exceeds 500,000 characters")
            rows = [
                (f"doc:line:{i}", line.strip()) for i, line in enumerate(text.splitlines()) if line.strip()
            ]
            resource.warnings.append("旧版 DOC 已转为文本；表格结构和条款关系仍需核对")
        except (subprocess.SubprocessError, UnicodeError) as error:
            resource.warnings.append("旧版 DOC 转换失败：" + type(error).__name__)
    elif suffix in {".txt", ".md"} or mime.startswith("text/plain"):
        rows = [
            (f"text:line:{i}", text.strip())
            for i, text in enumerate(raw.decode("utf-8").splitlines())
            if text.strip()
        ]
    else:
        resource.warnings.append("格式暂不支持；保留原文件，需人工转为 PDF/DOCX/XLSX/TXT")
    seed = digest(raw)[:16]
    resource.blocks = [
        {"id": f"{seed}:{i}", "locator": locator, "text": text} for i, (locator, text) in enumerate(rows)
    ]
    if not resource.blocks:
        resource.warnings.append("没有可提取文字")
    if sum(len(block["text"]) for block in resource.blocks) > 500000:
        raise AcquisitionError("extracted text exceeds 500,000 characters")
    return resource
