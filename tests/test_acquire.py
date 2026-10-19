import io
import socket
from contextlib import contextmanager

import httpx
import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from compute_policy.acquire import AcquisitionError, Fetcher, Resource, normalize, parse


def resolver(host, port, **kwargs):
    return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", port))]


def test_redirect_to_unregistered_host_is_rejected():
    calls = []

    def handle(request):
        calls.append(str(request.url))
        if request.url.path == "/robots.txt":
            return httpx.Response(404)
        return httpx.Response(302, headers={"Location": "http://169.254.169.254/latest/meta-data"})

    fetcher = Fetcher(["policy.example"], transport=httpx.MockTransport(handle), resolver=resolver, delay=0)
    with pytest.raises(AcquisitionError, match="allowlist"):
        fetcher.fetch("https://policy.example/notice")
    assert all("169.254" not in c for c in calls)


def test_public_host_with_private_dns_is_rejected():
    def private(host, port, **kwargs):
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1", port))]

    fetcher = Fetcher(["policy.example"], resolver=private)
    with pytest.raises(AcquisitionError, match="non-public"):
        fetcher.fetch("https://policy.example/notice")


def test_pinned_connections_prefer_ipv4_and_retry_only_validated_addresses(monkeypatch):
    def dual_stack(host, port, **kwargs):
        return [
            (socket.AF_INET6, socket.SOCK_STREAM, 6, "", ("2606:4700:4700::1111", port)),
            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("8.8.8.8", port)),
        ]

    calls = []
    fetcher = Fetcher(["policy.example"], resolver=dual_stack, delay=0)

    @contextmanager
    def stream(method, url, *, headers, extensions):
        calls.append(url.host)
        assert headers["Host"] == "policy.example"
        assert extensions["sni_hostname"] == "policy.example"
        if len(calls) == 1:
            raise httpx.ConnectError("unreachable address")
        yield httpx.Response(200, content=b"policy")

    monkeypatch.setattr(fetcher.client, "stream", stream)
    try:
        assert fetcher._request("https://policy.example/notice")[2] == b"policy"
        assert calls == ["8.8.8.8", "2606:4700:4700::1111"]
    finally:
        fetcher.close()


def test_robots_denial_is_not_empty_policy():
    fetcher = Fetcher(
        ["policy.example"],
        resolver=resolver,
        delay=0,
        transport=httpx.MockTransport(lambda request: httpx.Response(200, text="User-agent: *\nDisallow: /")),
    )
    with pytest.raises(AcquisitionError, match="disallows"):
        fetcher.fetch("https://policy.example/notice")


@pytest.mark.parametrize(
    "url", ["file:///etc/passwd", "http://name:pass@example.com", "http://example.com:8080"]
)
def test_bad_urls(url):
    with pytest.raises(ValueError):
        normalize(url)


def test_office_tables_and_blank_pdf_preserve_locations():
    document = Document()
    document.add_paragraph("政策正文要求企业提供合规资料。")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表格中的资格条件"
    data = io.BytesIO()
    document.save(data)
    parsed = parse(Resource("https://example.com/r.docx", data.getvalue(), "application/octet-stream"))
    assert any("table:0" in b["locator"] for b in parsed.blocks)
    workbook = Workbook()
    workbook.active["B2"] = "算力支出"
    data = io.BytesIO()
    workbook.save(data)
    parsed = parse(Resource("https://example.com/r.xlsx", data.getvalue(), "application/octet-stream"))
    assert parsed.blocks[0]["locator"].endswith(":B2")
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    data = io.BytesIO()
    writer.write(data)
    parsed = parse(Resource("https://example.com/r.pdf", data.getvalue(), "application/pdf"))
    assert any("OCR" in warning for warning in parsed.warnings)


def test_html_scripts_are_removed_and_attachment_is_resolved():
    parsed = parse(
        Resource(
            "https://example.com/a/index.html",
            b"<html><article><p>Policy rule</p>"
            b'<script>steal()</script><a href="../rule.pdf">Attachment</a></article></html>',
            "text/html",
        )
    )
    assert parsed.links[0]["url"] == "https://example.com/rule.pdf"
    assert parsed.links[0]["attachment"]
    assert not any("steal" in block["text"] for block in parsed.blocks)


def test_spreadsheet_formula_and_sparse_dimensions_fail_closed():
    workbook = Workbook()
    workbook.active["A1"] = "=SUM(B1:B5)"
    data = io.BytesIO()
    workbook.save(data)
    parsed = parse(Resource("https://example.com/a.xlsx", data.getvalue(), "application/octet-stream"))
    assert not parsed.blocks
    assert any("公式" in warning for warning in parsed.warnings)
    workbook.active["XFD1048576"] = "oversized dimensions"
    data = io.BytesIO()
    workbook.save(data)
    with pytest.raises(AcquisitionError, match="dimensions"):
        parse(Resource("https://example.com/a.xlsx", data.getvalue(), "application/octet-stream"))
